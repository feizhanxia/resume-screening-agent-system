from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT_ROOT = PROJECT_ROOT / "docs" / "assets" / "windows-guide"
OUTPUT = PROJECT_ROOT / "docs" / "简历筛选工作台-Windows版使用手册.docx"

SCREENSHOTS = {
    "settings": SCREENSHOT_ROOT / "01-settings.png",
    "task": SCREENSHOT_ROOT / "02-task-config.png",
    "criteria": SCREENSHOT_ROOT / "03-criteria.png",
    "progress": SCREENSHOT_ROOT / "04-progress.png",
    "recommended": SCREENSHOT_ROOT / "05-recommended.png",
    "uncertain": SCREENSHOT_ROOT / "06-uncertain.png",
}

TEAL = "103C3A"
TEAL_LIGHT = "E8F1EF"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667B80"
INK = "15353B"
GOLD = "D99342"
PALE_GOLD = "FFF4E5"
TABLE_FILL = "E8EEF5"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "1")


def set_font(run, *, size: float | None = None, bold: bool | None = None,
             color: str | None = None, italic: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    tail = paragraph.add_run(" 页")
    set_font(tail, size=9, color=MUTED)


def add_custom_numbering(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(level_text)
    alignment = OxmlElement("w:lvlJc")
    alignment.set(qn("w:val"), "left")
    level.append(alignment)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_element)


def add_numbered(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    apply_numbering(paragraph, num_id)
    set_font(paragraph.add_run(text), size=11, color=INK)


def add_bullet(doc: Document, text: str, num_id: int) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    apply_numbering(paragraph, num_id)
    set_font(paragraph.add_run(text), size=11, color=INK)


def add_callout(doc: Document, label: str, text: str, *, caution: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.14)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), PALE_GOLD if caution else TEAL_LIGHT)
    p_pr.append(shading)
    label_run = paragraph.add_run(f"{label}  ")
    set_font(label_run, size=10.5, bold=True, color=GOLD if caution else TEAL)
    set_font(paragraph.add_run(text), size=10.5, color=INK)


def add_caption(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(f"图 {number}  {text}")
    set_font(run, size=9.5, color=MUTED, italic=True)


def add_screenshot(doc: Document, key: str, number: int, caption: str) -> None:
    path = SCREENSHOTS[key]
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.45))
    inline = run._r.xpath(".//wp:inline")
    if inline:
        doc_pr = inline[0].find(qn("wp:docPr"))
        if doc_pr is not None:
            doc_pr.set("descr", caption)
    add_caption(doc, number, caption)


def add_heading(doc: Document, text: str, level: int = 1, *, new_page: bool = False) -> None:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.page_break_before = new_page
    set_keep_with_next(paragraph)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        set_font(prefix_run, size=11, bold=True, color=INK)
        text = text[len(bold_prefix):]
    set_font(paragraph.add_run(text), size=11, color=INK)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = paragraph.add_run("简历筛选工作台")
    set_font(left, size=9, bold=True, color=TEAL)
    right = paragraph.add_run("  ·  Windows 使用手册")
    set_font(right, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def build_manual() -> None:
    for path in SCREENSHOTS.values():
        if not path.exists():
            raise FileNotFoundError(path)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    decimal_num = add_custom_numbering(doc, bullet=False)
    bullet_num = add_custom_numbering(doc, bullet=True)

    # Editorial cover.
    doc.add_paragraph().paragraph_format.space_after = Pt(92)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("WINDOWS OPERATOR GUIDE"), size=10, bold=True, color=GOLD)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("简历筛选工作台"), size=30, bold=True, color=TEAL)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    set_font(subtitle.add_run("Windows x64 图文使用手册"), size=16, color=DARK_BLUE)
    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_after = Pt(108)
    set_font(lead.add_run("安装版与免安装版 · 本地优先 · 八路并行"), size=11, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(meta.add_run("适用版本 0.1.0  |  2026 年 8 月"), size=10, bold=True, color=TEAL)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(note.add_run("面向招聘人员和项目管理员"), size=9.5, italic=True, color=MUTED)
    add_heading(doc, "快速开始", 1, new_page=True)
    add_body(doc, "本工具用于批量扫描本地候选人材料，按岗位 JD 和动态筛选条件完成提取、分析、排序、Excel 导出及入选材料整理。原始 PDF、DOCX 和图片始终保存在本机；仅在需要语义判断时向配置的模型接口发送必要文本。")

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    mark_table_header(table.rows[0])
    headers = ("交付形式", "使用方式")
    for index, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], TABLE_FILL)
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=10.5, bold=True, color=INK)
    rows = [
        ("安装版", "双击 ResumeScreeningAgent-Setup-x64.exe，按向导完成安装，再从桌面或开始菜单启动。"),
        ("免安装版", "完整解压 ResumeScreeningAgent-portable-x64.zip，进入解压目录，双击 ResumeScreeningAgent.exe。"),
    ]
    for label, detail in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p0.add_run(label), size=10.5, bold=True, color=TEAL)
        set_font(cells[1].paragraphs[0].add_run(detail), size=10.5, color=INK)
    set_table_geometry(table, [2700, 6660])

    add_callout(doc, "免安装版注意", "不要直接在 ZIP 中运行，也不要只复制 EXE。必须保留完整的 _internal 文件夹；卸载时删除整个解压目录即可。", caution=True)

    add_heading(doc, "一次完整筛选流程", 2)
    for text in (
        "在“设置”中填写已轮换的新 DeepSeek API Key，并将并行 Worker 设为 8。",
        "选择简历根目录和输出目录，点击“扫描目录”核对候选人及文件数量。",
        "粘贴岗位 JD 和补充筛选要求，点击“重新解析筛选条件”。",
        "检查每条条件的启用状态、类型、权重和证据口径，再开始批量筛选。",
        "在处理进度页观察八路 Worker、完成数量和运行事件；需要时可暂停或取消。",
        "在候选人结果页核查证据、导出 Excel，并整理勾选候选人的完整材料目录。",
    ):
        add_numbered(doc, text, decimal_num)

    add_heading(doc, "1. 首次启动与模型设置", 1, new_page=True)
    add_screenshot(doc, "settings", 1, "设置 DeepSeek 凭据、模型与并行 Worker")
    add_body(doc, "点击左下角“设置”打开配置窗口。API Key 留空表示继续使用系统凭据；新填写的 Key 会写入 Windows 系统凭据库，不会进入项目、数据库、Excel 或日志。")
    for text in (
        "快速模型：用于条件解析和大多数候选人分析；填写账号实际可用的模型名称。",
        "复核模型：用于高价值或不确定结果的复核；可与快速模型相同。",
        "并行 Worker：推荐设置为 8。截图中的数值仅展示设置入口，不代表推荐默认值。",
        "隐私：建议启用发送前脱敏，减少电话、邮箱和身份证号进入模型请求。",
    ):
        add_bullet(doc, text, bullet_num)
    add_callout(doc, "安全要求", "不要继续使用曾经粘贴到聊天、邮件或公开文档中的旧 API Key；先在服务商控制台撤销并重新生成。", caution=True)
    add_heading(doc, "2. 新建筛选任务", 1, new_page=True)
    add_screenshot(doc, "task", 2, "选择目录、扫描候选人并输入岗位 JD")
    add_body(doc, "简历根目录采用“每位候选人一个一级子目录”的结构。候选人目录中的 PDF、DOCX、图片及其他补充材料会作为一个完整材料包登记；不要把不同候选人的材料混在同一子目录。")
    for text in (
        "简历根目录：选择包含全部候选人子目录的文件夹。",
        "输出目录：用于保存缓存、Excel、运行记录和整理后的候选人材料。建议与原始简历目录分开。",
        "扫描目录：确认候选人数、文件数、可提取文件数和未归属文件数；数量异常时先整理目录。",
        "岗位 JD：粘贴完整职责、任职要求和加分项，不必预先拆成固定字段。",
        "补充筛选要求：填写本次特有要求，例如优先项目类型、必须人工核查的院校规则或证书。",
    ):
        add_bullet(doc, text, bullet_num)
    add_callout(doc, "核对原则", "扫描到的候选人数应等于一级候选人子目录数；文件数应覆盖所有简历和补充材料。未归属文件不应直接忽略。")
    add_heading(doc, "3. 确认筛选条件", 1, new_page=True)
    add_screenshot(doc, "criteria", 3, "在启动前检查动态拆解出的筛选条件")
    add_body(doc, "系统会根据当前 JD 动态生成条件。表格可以直接编辑，也可以新增或删除条件。建议先确认“证据口径”，再调整类型和权重。")
    for text in (
        "硬性：不满足时会显著影响推荐，但证据不足时应保留“不确定”，不能自动判定为不符合。",
        "评分：用于拉开候选人差异，例如技术深度、项目质量和跨学科协作。",
        "加分：用于排序接近者，例如论文、竞赛、HPC、CUDA 或特定行业经验。",
        "需外部核验：仅在确实需要院校名单、排名或其他外部事实时启用；无法获得可信依据时保持“不确定”。",
    ):
        add_bullet(doc, text, bullet_num)
    add_callout(doc, "避免误筛", "985/211、双一流、排名前 20 等口径不能只凭模型记忆判断。没有明确年份、名单和证据时，应进入人工复核，而不是直接淘汰。", caution=True)
    add_heading(doc, "4. 八路并行处理", 1, new_page=True)
    add_screenshot(doc, "progress", 4, "查看总体进度、Worker 状态和运行事件")
    add_body(doc, "八名 Worker 动态领取候选人；同一候选人的全部材料始终由一个 Worker 合并分析。进度页显示完成数量、当前处理对象和最近事件。")
    for text in (
        "暂停：停止领取新任务；已开始的单个提取或模型调用会安全结束。",
        "继续：从现有队列恢复，不重复处理已经完成且缓存有效的材料。",
        "取消：终止当前运行；已写入缓存和数据库的终态会保留，后续可增量继续。",
        "异常材料：加密、空白、图片型或解析失败文件会保留记录并标记人工复核，不应静默消失。",
    ):
        add_bullet(doc, text, bullet_num)
    add_callout(doc, "性能建议", "默认使用 8 路。继续提高并发不一定更快，可能触发模型限流、磁盘争用或 OCR 内存压力。先用 20–50 份材料试跑，再处理完整目录。")
    add_heading(doc, "5. 查看推荐结果与证据", 1, new_page=True)
    add_screenshot(doc, "recommended", 5, "推荐候选人的条件结论和原文证据")
    add_body(doc, "结果默认按综合分排序。点击任意候选人，右侧会显示总分、建议、摘要，以及每条启用条件的判断、理由和原文证据位置。")
    for text in (
        "符合：证据明确支持该条件。",
        "不符合：材料中存在明确反证，或明确未达到要求。",
        "不确定：材料未写明、文本提取不足、条件需要外部事实，或模型不能可靠确认。",
        "待复核：整体结果可供排序，但至少有一项关键条件需要招聘人员确认。",
    ):
        add_bullet(doc, text, bullet_num)
    add_callout(doc, "证据优先", "分数用于排序，不替代人工决定。重点候选人应检查右侧证据是否来自正确文件和页码，尤其是学历、任职时间和核心项目经历。")
    add_heading(doc, "6. 处理不确定与接近匹配", 1, new_page=True)
    add_screenshot(doc, "uncertain", 6, "关键条件证据不足时保留不确定并提示人工复核")
    add_body(doc, "复杂岗位通常没有完全符合的人选。系统会保留最接近的候选人，并把缺口拆到具体条件中；招聘人员可以在不重新提取简历的情况下调整条件或权重后重新分析。")
    review_num = add_custom_numbering(doc, bullet=False)
    for text in (
        "先检查“不确定”是否来自简历未写明，而不是确实不满足。",
        "对排名、院校类别、论文层级等外部事实，补充可信依据后再复核。",
        "对同义技能和相邻领域经验，查看原文项目证据，不要只看关键词是否完全一致。",
        "如果修改 JD 或条件，仅重算受影响的判断；文件内容未变化时复用提取缓存。",
    ):
        add_numbered(doc, text, review_num)
    add_heading(doc, "7. 导出 Excel 与整理候选人材料", 1, new_page=True)
    add_body(doc, "候选人结果页顶部提供搜索、最低分过滤、“导出 Excel”和“整理勾选材料”。建议先筛选并逐项核查证据，再进行导出和材料整理。")
    export_num = add_custom_numbering(doc, bullet=False)
    for text in (
        "搜索：按姓名、最近公司或技能关键词缩小结果范围。",
        "最低分：隐藏低于阈值的候选人，但不会删除数据库记录。",
        "导出 Excel：生成候选人结果、条件明细、证据、文件清单、覆盖审计等工作表；HR 状态、负责人和备注可人工维护。",
        "整理勾选材料：勾选候选人后复制其完整子目录，保留相对路径并执行文件哈希校验。",
        "再次运行：新增或修改材料会重新处理；未变化材料和已提取文本会复用。",
    ):
        add_numbered(doc, text, export_num)

    add_heading(doc, "8. 常见问题", 1, new_page=True)
    faq = [
        ("程序无法启动", "确认 ZIP 已完整解压，ResumeScreeningAgent.exe 与 _internal 文件夹处于同一级目录。不要只复制 EXE。"),
        ("Windows SmartScreen 提示", "当前安装器尚未使用商业代码签名证书。先核对文件 SHA-256；确认来源后可选择“更多信息 → 仍要运行”。"),
        ("PDF 没有文字", "图片型页面会尝试内置离线 OCR。仍失败时保留失败终态和文件记录，转人工查看。"),
        ("运行速度没有明显提升", "检查模型接口限流、网络延迟、OCR 比例和磁盘性能。并发高于 8 不保证更快。"),
        ("新增了简历", "重新选择或扫描同一根目录；新增和变化文件进入队列，缓存有效的旧文件不会重复提取。"),
        ("需要修改筛选口径", "返回任务配置或条件页修改 JD、补充要求、条件类型或权重，再启动分析；通常复用已有文本缓存。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    mark_table_header(table.rows[0])
    for index, text in enumerate(("问题", "处理方法")):
        set_cell_shading(table.rows[0].cells[index], TABLE_FILL)
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=10.5, bold=True, color=INK)
    for question, answer in faq:
        cells = table.add_row().cells
        set_font(cells[0].paragraphs[0].add_run(question), size=10, bold=True, color=TEAL)
        set_font(cells[1].paragraphs[0].add_run(answer), size=10, color=INK)
    set_table_geometry(table, [2500, 6860])

    add_callout(doc, "发布提示", "对外分发前建议为安装器配置组织自己的 Windows 代码签名证书，并使用真实的 500 份简历样本完成性能、异常文件和覆盖率验收。", caution=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "简历筛选工作台 Windows 版使用手册"
    doc.core_properties.subject = "安装、配置、批量筛选、结果复核和材料整理"
    doc.core_properties.author = "Resume Screening Team"
    doc.core_properties.keywords = "简历筛选, Windows, 使用手册, 八路并行"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()
