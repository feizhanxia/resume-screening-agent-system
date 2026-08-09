from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image

from resume_agent.audit.coverage import AuditReport
from resume_agent.domain.models import (
    AnalysisResult,
    CandidateFacts,
    CandidatePackage,
    Criterion,
    Evidence,
    ExtractedUnit,
    ExtractionResult,
    Judgment,
    MaterialFile,
)
from resume_agent.extraction import docx as docx_extractor
from resume_agent.extraction import image as image_extractor
from resume_agent.llm.deepseek import DeepSeekClient, ModelAnalysisPayload
from resume_agent.materials.copying import copy_candidate_packages, safe_relative_path
from resume_agent.reporting.excel import export_excel
from resume_agent.safety import escape_rich_text, excel_safe_value
from resume_agent.scanning.scanner import scan_resume_root, sha256_file


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)", "  =1+1", "\t=1+1"])
def test_excel_safe_value_neutralizes_formula_prefixes(value: str):
    assert excel_safe_value(value).startswith("'")


def test_excel_export_writes_untrusted_text_as_plain_text(tmp_path: Path):
    root = tmp_path / "resumes"
    folder = root / "=malicious"
    folder.mkdir(parents=True)
    material = folder / "resume.txt"
    material.write_text("safe", encoding="utf-8")
    scan = scan_resume_root(root)
    criterion = Criterion(criterion_id="c1", name="=name", definition="=definition")
    analysis = AnalysisResult(
        candidate_id=scan.candidates[0].candidate_id,
        facts=CandidateFacts(name="=HYPERLINK(\"https://invalid\")", summary="+payload"),
        judgments=[
            Judgment(
                criterion_id="c1",
                conclusion="uncertain",
                confidence=0,
                reason="@payload",
            )
        ],
        overall_reason="-payload",
    )
    output = export_excel(
        tmp_path / "result.xlsx",
        "run_test",
        scan,
        [criterion],
        [analysis],
        AuditReport(
            passed=True,
            candidate_count=1,
            file_count=1,
            analysis_count=1,
        ),
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook["候选人信息"]["B2"].data_type == "s"
    assert workbook["候选人信息"]["B2"].value.startswith("'")
    assert workbook["判断明细"]["E2"].data_type == "s"
    assert workbook["当前结果"]["H2"].data_type == "s"


def test_scanner_rejects_symlinked_material(tmp_path: Path):
    root = tmp_path / "resumes"
    folder = root / "candidate"
    folder.mkdir(parents=True)
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    link = folder / "resume.txt"
    try:
        link.symlink_to(outside)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks are unavailable on this platform")

    result = scan_resume_root(root)

    assert result.file_count == 0
    assert any("符号链接" in error or "Junction" in error for error in result.errors)


def test_material_copy_revalidates_source_and_relative_path(tmp_path: Path):
    root = tmp_path / "candidate"
    root.mkdir()
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    digest = sha256_file(outside)
    package = CandidatePackage(
        candidate_id="cand_test",
        display_name="candidate",
        root_path=root,
        relative_dir="candidate",
        package_hash=digest,
        files=[
            MaterialFile(
                file_id="file_test",
                candidate_id="cand_test",
                absolute_path=outside,
                relative_path="resume.txt",
                media_type="text/plain",
                content_hash=digest,
                size=outside.stat().st_size,
                mtime_ns=outside.stat().st_mtime_ns,
            )
        ],
    )

    with pytest.raises(ValueError, match="候选人目录"):
        copy_candidate_packages([package], {"cand_test"}, tmp_path / "output", "run")
    with pytest.raises(ValueError):
        safe_relative_path("../escape.txt")
    with pytest.raises(ValueError):
        safe_relative_path("C:\\escape.txt")


def test_rich_text_is_html_escaped():
    assert escape_rich_text('<img src="file:///secret">') == (
        "&lt;img src=&quot;file:///secret&quot;&gt;"
    )


def test_model_document_chunking_preserves_all_text_and_hides_local_path(tmp_path: Path):
    material = tmp_path / "张三简历.txt"
    material.write_text("abcdefghijk", encoding="utf-8")
    digest = sha256_file(material)
    file = MaterialFile(
        file_id="file_abc",
        candidate_id="cand_test",
        absolute_path=material,
        relative_path="张三简历.txt",
        media_type="text/plain",
        content_hash=digest,
        size=material.stat().st_size,
        mtime_ns=material.stat().st_mtime_ns,
    )
    package = CandidatePackage(
        candidate_id="cand_test",
        display_name="张三",
        root_path=tmp_path,
        relative_dir="张三",
        package_hash=digest,
        files=[file],
    )
    extraction = ExtractionResult(
        file_id=file.file_id,
        status="extracted",
        units=[
            ExtractedUnit(
                locator_kind="page",
                locator_value="1",
                text="abcdefghijk",
                method="native",
            )
        ],
    )

    chunks = DeepSeekClient._build_document_chunks(
        package, {file.file_id: extraction}, True, max_chars=5
    )
    sent = "".join(
        unit["text"]
        for chunk in chunks
        for document in chunk
        for unit in document["units"]
    )

    assert sent == "abcdefghijk"
    assert len(chunks) == 3
    assert all(document["relative_path"] == "file_abc" for chunk in chunks for document in chunk)


def test_model_evidence_must_match_local_locator_and_quote():
    criterion = Criterion(criterion_id="c1", name="Python", definition="熟悉 Python")
    parsed = ModelAnalysisPayload(
        judgments=[
            Judgment(
                criterion_id="c1",
                conclusion="match",
                confidence=0.99,
                reason="符合",
                evidence=[
                    Evidence(
                        file_id="file_abc",
                        relative_path="file_abc",
                        locator_kind="page",
                        locator_value="1",
                        quote="伪造原文",
                    )
                ],
            )
        ]
    )

    judgments = DeepSeekClient._validated_judgments(
        parsed,
        [criterion],
        {"file_abc": "resume.pdf"},
        {("file_abc", "page", "1"): "真实原文：熟悉 Python"},
    )

    assert judgments[0].conclusion == "uncertain"
    assert judgments[0].evidence == []


def test_docx_archive_limit_fails_visibly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Python engineer")
    document.save(path)
    monkeypatch.setattr(docx_extractor, "MAX_ARCHIVE_ENTRIES", 1)

    result = docx_extractor.extract_docx(path, "file_docx")

    assert result.status == "failed"
    assert "安全上限" in (result.error or "")


def test_image_pixel_limit_fails_visibly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "resume.png"
    Image.new("RGB", (10, 10), "white").save(path)
    monkeypatch.setattr(image_extractor, "MAX_IMAGE_PIXELS", 50)

    result = image_extractor.extract_image(path, "file_image")

    assert result.status == "failed"
    assert "像素" in (result.error or "")
