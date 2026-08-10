import asyncio
import json
import os
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import load_workbook

from resume_agent.audit.coverage import audit_run
from resume_agent.config import WorkspacePaths
from resume_agent.criteria.compiler import heuristic_compile
from resume_agent.domain.models import (
    CandidatePackage,
    ExtractedUnit,
    ExtractionResult,
    MaterialFile,
    RunConfig,
)
from resume_agent.extraction.service import ExtractionService
from resume_agent.llm.deepseek import (
    DeepSeekClient,
    MaterialRoutingOutcome,
    ModelCallError,
)
from resume_agent.reporting.excel import export_excel
from resume_agent.scanning.scanner import classify_role, scan_resume_root
from resume_agent.storage.database import Database
from resume_agent.workers.engine import ScreeningEngine, effective_worker_count


def _material(tmp_path: Path, file_id: str, name: str, role: str) -> MaterialFile:
    path = tmp_path / name
    path.write_text("fixture", encoding="utf-8")
    return MaterialFile(
        file_id=file_id,
        candidate_id="cand_test",
        absolute_path=path,
        relative_path=name,
        media_type="application/pdf",
        material_role=role,
        content_hash=file_id,
        size=path.stat().st_size,
        mtime_ns=path.stat().st_mtime_ns,
    )


def _extraction(file_id: str, pages: list[str]) -> ExtractionResult:
    return ExtractionResult(
        file_id=file_id,
        status="extracted",
        expected_units=len(pages),
        covered_units=len(pages),
        units=[
            ExtractedUnit(
                locator_kind="page",
                locator_value=str(index),
                text=text,
                method="native",
            )
            for index, text in enumerate(pages, 1)
        ],
    )


def test_cv_filename_matching_does_not_treat_cvpr_as_resume(tmp_path: Path):
    assert classify_role(tmp_path / "candidate-CV.pdf") == "primary_resume"
    assert classify_role(tmp_path / "CVPR2026-paper.pdf") == "publication"


def test_routed_context_deduplicates_and_caps_non_resume_material(tmp_path: Path):
    resume = _material(tmp_path, "file_resume", "简历.pdf", "primary_resume")
    duplicate = resume.model_copy(update={"relative_path": "简历-副本.pdf"})
    publication = _material(tmp_path, "file_paper", "paper.pdf", "publication")
    package = CandidatePackage(
        candidate_id="cand_test",
        display_name="候选人",
        root_path=tmp_path,
        relative_dir="候选人",
        package_hash="hash",
        files=[resume, duplicate, publication],
    )
    extractions = {
        "file_resume": _extraction("file_resume", ["R" * 9_000]),
        "file_paper": _extraction("file_paper", ["A" * 5_000, "B" * 5_000, "C" * 5_000]),
    }

    chunks = DeepSeekClient._build_document_chunks(
        package, extractions, redact=False, max_chars=200_000
    )
    sent: dict[str, str] = {}
    locators: dict[str, set[str]] = {}
    for document in chunks[0]:
        sent.setdefault(document["file_id"], "")
        locators.setdefault(document["file_id"], set())
        for unit in document["units"]:
            sent[document["file_id"]] += unit["text"]
            locators[document["file_id"]].add(unit["locator_value"])

    assert len(sent["file_resume"]) == 9_000
    assert len(sent["file_paper"]) == 8_000
    assert locators["file_paper"] == {"1", "2"}


def test_json_chat_retries_and_disables_thinking(monkeypatch):
    client = DeepSeekClient("sk-test", RunConfig(max_retries_per_resume=2))
    responses = iter(["", "not-json", '{"ok": true}'])
    requests = []

    async def create(**kwargs):
        requests.append(kwargs)
        return SimpleNamespace(
            choices=[
                SimpleNamespace(
                    finish_reason="stop",
                    message=SimpleNamespace(content=next(responses)),
                )
            ]
        )

    async def no_sleep(_seconds):
        return None

    client.client = SimpleNamespace(
        chat=SimpleNamespace(completions=SimpleNamespace(create=create))
    )
    monkeypatch.setattr("resume_agent.llm.deepseek.asyncio.sleep", no_sleep)

    payload = asyncio.run(client._json_chat(model="test", system="json", user="{}"))

    assert payload == {"ok": True}
    assert len(requests) == 3
    assert all(
        request["extra_body"] == {"thinking": {"type": "disabled"}}
        for request in requests
    )


def test_material_preclassification_uses_bounded_previews(tmp_path: Path):
    resume = _material(tmp_path, "file_resume", "材料一.pdf", "supplementary")
    paper = _material(tmp_path, "file_paper", "材料二.pdf", "supplementary")
    package = CandidatePackage(
        candidate_id="cand_test",
        display_name="候选人",
        root_path=tmp_path,
        relative_dir="候选人",
        package_hash="hash",
        files=[resume, paper],
    )
    extractions = {
        "file_resume": _extraction("file_resume", ["教育经历 工作经历 " + "R" * 10_000]),
        "file_paper": _extraction("file_paper", ["Abstract " + "P" * 10_000]),
    }
    client = DeepSeekClient("sk-test", RunConfig(redact_pii=False))
    captured = {}

    async def route_json(**kwargs):
        captured.update(json.loads(kwargs["user"]))
        return {
            "materials": [
                {"file_id": "file_resume", "role": "primary_resume", "confidence": 0.9},
                {"file_id": "file_paper", "role": "publication", "confidence": 0.9},
            ]
        }

    client._json_chat = route_json
    outcome = asyncio.run(client.classify_materials(package, extractions))

    previews = [item["preview"] for item in captured["materials"]]
    assert all(len(preview) <= 4_000 for preview in previews)
    assert sum(map(len, previews)) <= 30_000
    assert outcome.roles == {
        "file_resume": "primary_resume",
        "file_paper": "publication",
    }
    assert not outcome.needs_review


def test_concurrent_duplicate_extraction_is_coalesced(tmp_path: Path, monkeypatch):
    root = tmp_path / "resumes"
    candidate = root / "候选人"
    candidate.mkdir(parents=True)
    document = Document()
    document.add_paragraph("Python")
    document.save(candidate / "简历.docx")
    shutil.copyfile(candidate / "简历.docx", candidate / "简历-副本.docx")
    scan = scan_resume_root(root)
    database = Database(tmp_path / "cache.db")
    database.save_scan(scan.candidates)
    service = ExtractionService(database, concurrency=8)
    calls = 0

    from resume_agent.extraction import service as service_module

    original = service_module.extract_docx

    def counted(path, file_id):
        nonlocal calls
        calls += 1
        return original(path, file_id)

    monkeypatch.setattr(service_module, "extract_docx", counted)

    async def extract_all():
        return await asyncio.gather(
            *(service.extract(file) for file in scan.candidates[0].files)
        )

    asyncio.run(extract_all())

    assert calls == 1


def test_duplicate_files_extract_once_and_model_failure_degrades(tmp_path: Path):
    root = tmp_path / "resumes"
    candidate = root / "候选人"
    candidate.mkdir(parents=True)
    document = Document()
    document.add_paragraph("教育经历 工作经历 Python 项目经验")
    document.save(candidate / "简历.docx")
    shutil.copyfile(candidate / "简历.docx", candidate / "简历-副本.docx")
    output = tmp_path / "output"
    workspace = WorkspacePaths.create(output)
    database = Database(workspace.database)
    scan = scan_resume_root(root)
    criteria = heuristic_compile("要求熟悉 Python")
    engine = ScreeningEngine(
        database,
        RunConfig(worker_count=10, max_retries_per_resume=2),
        api_key="sk-test",
    )
    extraction_calls = 0
    original_extract = engine.extractions.extract

    async def counted_extract(file):
        nonlocal extraction_calls
        extraction_calls += 1
        return await original_extract(file)

    async def route(_package, _extractions):
        return MaterialRoutingOutcome(
            roles={scan.candidates[0].files[0].file_id: "primary_resume"}
        )

    async def fail_analysis(_package, _extractions, _criteria):
        raise ModelCallError(
            {
                "error_type": "JSONDecodeError",
                "finish_reason": "stop",
                "response_chars": 0,
                "attempts": 3,
            }
        )

    engine.extractions.extract = counted_extract
    engine.llm.classify_materials = route
    engine.llm.analyze = fail_analysis

    run_id, analyses = asyncio.run(engine.run(scan, criteria, "要求熟悉 Python", "", output))

    assert extraction_calls == 1
    assert scan.file_count == 2
    assert len(analyses) == 1
    assert analyses[0].recommendation == "review"
    assert analyses[0].prompt_version == "routing-v1-degraded"
    audit = audit_run(scan, criteria, analyses, database)
    excel = export_excel(
        workspace.outputs / "duplicate-result.xlsx",
        run_id,
        scan,
        criteria,
        analyses,
        audit,
    )
    workbook = load_workbook(excel, read_only=True)
    material_paths = {
        workbook["文件清单"].cell(row, 4).value
        for row in range(2, workbook["文件清单"].max_row + 1)
    }
    assert material_paths == {"简历.docx", "简历-副本.docx"}
    with database.connection() as conn:
        status = conn.execute("SELECT status FROM runs WHERE run_id=?", (run_id,)).fetchone()[0]
        event = conn.execute(
            "SELECT payload_json FROM events WHERE run_id=? AND event_type='candidate_degraded'",
            (run_id,),
        ).fetchone()[0]
    assert status == "completed_with_errors"
    assert json.loads(event)["error_type"] == "JSONDecodeError"
    assert "prompt" not in event.lower()


def test_effective_worker_count_uses_configuration_and_candidate_limit():
    assert effective_worker_count(10, 15) == 10
    assert effective_worker_count(10, 4) == 4
    assert effective_worker_count(10, 0) == 0


def test_worker_labels_are_rebuilt_to_actual_count(monkeypatch, tmp_path: Path):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    try:
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QApplication
    except ImportError as exc:
        pytest.skip(f"Qt runtime is unavailable: {exc}")

    from resume_agent.app.window import MainWindow

    app = QApplication.instance() or QApplication([])
    settings = QSettings(str(tmp_path / "settings.ini"), QSettings.Format.IniFormat)
    monkeypatch.setattr("resume_agent.app.window.QSettings", lambda: settings)
    window = MainWindow()
    window._reset_worker_labels(10)

    assert len(window.worker_labels) == 10
    assert window.worker_labels[-1].text().startswith("Worker 10")
    window.close()
    app.processEvents()
