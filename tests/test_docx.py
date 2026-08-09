from pathlib import Path

from docx import Document

from resume_agent.extraction.docx import extract_docx


def test_docx_preserves_paragraph_and_table_evidence(tmp_path: Path):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("张三 Python 工程师")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "PyTorch"
    document.save(path)

    result = extract_docx(path, "file_test")

    assert result.status == "extracted"
    assert [unit.locator_kind for unit in result.units] == ["paragraph", "table"]
    assert "PyTorch" in result.text
