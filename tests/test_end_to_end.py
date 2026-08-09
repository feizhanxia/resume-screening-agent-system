import asyncio
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from resume_agent.audit.coverage import audit_run
from resume_agent.config import WorkspacePaths
from resume_agent.criteria.compiler import heuristic_compile
from resume_agent.domain.models import RunConfig
from resume_agent.materials.copying import copy_candidate_packages
from resume_agent.reporting.excel import export_excel
from resume_agent.scanning.scanner import scan_resume_root
from resume_agent.storage.database import Database
from resume_agent.workers.engine import ScreeningEngine


def test_offline_pipeline_exports_and_copies(tmp_path: Path):
    root = tmp_path / "resumes"
    candidate = root / "张三"
    candidate.mkdir(parents=True)
    document = Document()
    document.add_paragraph("张三，硕士，Python 与 PyTorch 项目经验")
    document.save(candidate / "简历.docx")
    (candidate / "备注.txt").write_text("补充材料", encoding="utf-8")
    output = tmp_path / "output"
    workspace = WorkspacePaths.create(output)
    database = Database(workspace.database)
    scan = scan_resume_root(root)
    criteria = heuristic_compile("要求硕士及以上，精通 Python")
    engine = ScreeningEngine(database, RunConfig(worker_count=8), api_key=None)

    run_id, analyses = asyncio.run(engine.run(scan, criteria, "要求硕士", "", output))

    assert len(analyses) == 1
    assert len(analyses[0].judgments) == len(criteria)
    assert analyses[0].model == "offline"
    audit = audit_run(scan, criteria, analyses, database)
    assert not audit.passed
    database.save_manual_fields(analyses[0].candidate_id, "待面试", "HR-A", "电话沟通")
    excel = export_excel(
        workspace.outputs / "result.xlsx",
        run_id,
        scan,
        criteria,
        analyses,
        audit,
        database.all_manual_fields(),
    )
    assert set(load_workbook(excel, read_only=True).sheetnames) == {
        "运行状态",
        "文件清单",
        "候选人信息",
        "筛选条件",
        "判断明细",
        "当前结果",
        "未归属与异常",
    }
    workbook = load_workbook(excel, read_only=True)
    assert [workbook["当前结果"].cell(2, column).value for column in (9, 10, 11)] == [
        "待面试",
        "HR-A",
        "电话沟通",
    ]
    manifest = copy_candidate_packages(
        scan.candidates,
        {scan.candidates[0].candidate_id},
        workspace.outputs,
        run_id,
    )
    assert manifest.exists()
    assert list(manifest.parent.rglob("简历.docx"))


def test_more_than_bounded_queue_capacity_is_processed(tmp_path: Path):
    root = tmp_path / "bulk"
    for index in range(60):
        folder = root / f"candidate-{index:03d}"
        folder.mkdir(parents=True)
        (folder / "material.txt").write_text(str(index), encoding="utf-8")
    output = tmp_path / "output"
    workspace = WorkspacePaths.create(output)
    database = Database(workspace.database)
    scan = scan_resume_root(root)
    criteria = heuristic_compile("要求熟悉 Python")
    engine = ScreeningEngine(database, RunConfig(worker_count=8), api_key=None)

    _, analyses = asyncio.run(engine.run(scan, criteria, "要求熟悉 Python", "", output))

    assert len(analyses) == 60
