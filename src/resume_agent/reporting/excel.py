from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from resume_agent.audit.coverage import AuditReport
from resume_agent.domain.models import AnalysisResult, Criterion, ScanResult

HEADER_FILL = PatternFill("solid", fgColor="17324D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="E6F4F1")


def _write_header(sheet, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _fit(sheet, max_width: int = 48) -> None:
    for column_cells in sheet.columns:
        width = min(
            max(len(str(cell.value or "")) for cell in column_cells) + 2,
            max_width,
        )
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(10, width)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_excel(
    path: Path,
    run_id: str,
    scan: ScanResult,
    criteria: list[Criterion],
    analyses: list[AnalysisResult],
    audit: AuditReport,
    manual_fields: dict[str, dict[str, str]] | None = None,
) -> Path:
    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    status = workbook.active
    status.title = "运行状态"
    _write_header(status, ["项目", "值"])
    status_rows = [
        ("运行ID", run_id),
        ("导出时间", datetime.now().isoformat(timespec="seconds")),
        ("简历根目录", str(scan.root_path)),
        ("候选人数", len(scan.candidates)),
        ("文件数", scan.file_count),
        ("分析结果数", len(analyses)),
        ("未归属文件数", len(scan.unassigned_files)),
        ("覆盖审计", "通过" if audit.passed else "需复核"),
        ("审计发现", "；".join(audit.findings)),
    ]
    for row in status_rows:
        status.append(row)

    manifest = workbook.create_sheet("文件清单")
    _write_header(
        manifest,
        ["候选人ID", "候选人目录", "文件ID", "相对路径", "类型", "材料角色", "大小", "哈希"],
    )
    for candidate in scan.candidates:
        for file in candidate.files:
            manifest.append(
                [
                    candidate.candidate_id,
                    candidate.display_name,
                    file.file_id,
                    file.relative_path,
                    file.media_type,
                    file.material_role,
                    file.size,
                    file.content_hash,
                ]
            )

    candidate_sheet = workbook.create_sheet("候选人信息")
    _write_header(
        candidate_sheet,
        ["候选人ID", "姓名", "最高学历", "本科院校", "最近公司", "工作年限", "技能", "摘要"],
    )
    by_candidate = {item.candidate_id: item for item in analyses}
    for candidate in scan.candidates:
        result = by_candidate.get(candidate.candidate_id)
        facts = result.facts if result else None
        candidate_sheet.append(
            [
                candidate.candidate_id,
                (facts.name if facts and facts.name else candidate.display_name),
                facts.highest_degree if facts else "",
                facts.undergraduate_school if facts else "",
                facts.current_or_last_company if facts else "",
                facts.work_years if facts else "",
                "、".join(facts.skills) if facts else "",
                facts.summary if facts else "",
            ]
        )

    criteria_sheet = workbook.create_sheet("筛选条件")
    _write_header(criteria_sheet, ["条件ID", "名称", "定义", "类型", "权重", "启用", "外部事实"])
    for item in criteria:
        criteria_sheet.append(
            [
                item.criterion_id,
                item.name,
                item.definition,
                item.kind,
                item.weight,
                "是" if item.enabled else "否",
                "是" if item.requires_external_fact else "否",
            ]
        )

    detail = workbook.create_sheet("判断明细")
    _write_header(
        detail,
        ["候选人ID", "条件ID", "结论", "置信度", "理由", "证据位置", "证据原文", "待外部核验"],
    )
    for result in analyses:
        for judgment in result.judgments:
            locators = "\n".join(
                f"{e.relative_path} · {e.locator_kind} {e.locator_value}" for e in judgment.evidence
            )
            quotes = "\n".join(e.quote for e in judgment.evidence)
            detail.append(
                [
                    result.candidate_id,
                    judgment.criterion_id,
                    judgment.conclusion,
                    judgment.confidence,
                    judgment.reason,
                    locators,
                    quotes,
                    "是" if judgment.needs_external_fact else "否",
                ]
            )

    current = workbook.create_sheet("当前结果")
    _write_header(
        current,
        [
            "排名",
            "候选人ID",
            "姓名",
            "总分",
            "建议",
            "硬性项",
            "不确定项",
            "综合说明",
            "HR状态",
            "负责人",
            "HR备注",
        ],
    )
    criterion_by_id = {item.criterion_id: item for item in criteria}
    manual_fields = manual_fields or {}
    for rank, result in enumerate(sorted(analyses, key=lambda item: item.score, reverse=True), 1):
        required = [
            item.conclusion
            for item in result.judgments
            if criterion_by_id.get(item.criterion_id)
            and criterion_by_id[item.criterion_id].kind == "required"
        ]
        uncertain = sum(item.conclusion == "uncertain" for item in result.judgments)
        manual = manual_fields.get(result.candidate_id, {})
        current.append(
            [
                rank,
                result.candidate_id,
                result.facts.name or "",
                result.score,
                result.recommendation,
                "通过"
                if required and all(value == "match" for value in required)
                else "存在缺口/待核验",
                uncertain,
                result.overall_reason,
                manual.get("hr_status", ""),
                manual.get("owner", ""),
                manual.get("note", ""),
            ]
        )
    for row in current.iter_rows(min_row=2, max_col=4):
        row[3].fill = ACCENT_FILL

    unassigned = workbook.create_sheet("未归属与异常")
    _write_header(unassigned, ["类型", "路径或信息"])
    for file in scan.unassigned_files:
        unassigned.append(["未归属文件", str(file)])
    for error in scan.errors:
        unassigned.append(["扫描错误", error])
    for finding in audit.findings:
        unassigned.append(["审计发现", finding])

    for sheet in workbook.worksheets:
        _fit(sheet)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False, dir=path.parent) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)
    return path
